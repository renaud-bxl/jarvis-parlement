import base64
import json
import plistlib
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, Mock
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import common,services,providers,templates,desktop,chatgpt_auth
from docx import Document

class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.home=Path(self.temp.name);self.cfg=self.home/'config';self.cfg.mkdir()
        self.work=self.home/'docs';self.base=self.home/'app';self.base.mkdir()
        for module in [common,services,providers,templates,desktop,chatgpt_auth]:
            for name,value in [('HOME',self.home),('CONFIG',self.cfg),('WORK',self.work),('REPORTS',self.work/'08-Rapports'),('BASE',self.base)]:
                if hasattr(module,name):
                    p=patch.object(module,name,value);p.start();self.addCleanup(p.stop)
        p=patch.object(templates,'DIR',self.cfg/'templates');p.start();self.addCleanup(p.stop)
        common.initialize()
        services.save_settings({'full_name':'Camille Exemple','role':'Députée','institution':'Parlement bruxellois','model':'local:test'})
    def test_fresh_profile_has_no_inherited_name(self):
        (self.cfg/'settings.json').unlink();common.initialize()
        self.assertEqual(common.settings()['full_name'],'')
        with self.assertRaises(ValueError):services.instructions()
    def test_migration_preserves_account_model_and_schedule(self):
        common.atomic_json(self.cfg/'settings.json',{'model':'existing:test','custom':'preserved'})
        common.atomic_json(self.cfg/'accounts.json',{'existing':{'type':'gmail','address':'test@example.org'}})
        p=self.home/'Library/LaunchAgents/net.compose.jarvis-parlement.briefing.plist';p.parent.mkdir(parents=True)
        p.write_bytes(plistlib.dumps({'StartCalendarInterval':{'Hour':10,'Minute':30}}))
        common.initialize();cfg=common.settings()
        self.assertEqual(cfg['model'],'existing:test');self.assertEqual(cfg['local_model'],'existing:test')
        self.assertTrue(cfg['schedule_enabled']);self.assertEqual(cfg['briefing_time'],'10:30')
        self.assertEqual(services.accounts()['existing']['type'],'gmail');self.assertEqual(cfg['custom'],'preserved')
    def test_private_document_never_searches_web(self):
        f={'name':'private.txt','content':base64.b64encode(b'Secret').decode()}
        with patch.object(services,'complete') as call:
            with self.assertRaises(ValueError):services.chat({'query':'analyse','document':f,'web':True})
            call.assert_not_called()
    def test_private_history_never_searches_web(self):
        with patch.object(services,'complete') as call:
            with self.assertRaises(ValueError):services.chat({'query':'continue','web':True,'history':[{'private':True,'text':'confidentiel'}]})
            call.assert_not_called()
    def test_cloud_private_consent_is_required(self):
        for provider in ('chatgpt','openai','anthropic'):
            with self.assertRaises(ValueError):providers.complete('secret',{'provider':provider,'model':'test'},private=True)
    def test_cloud_ollama_is_rejected(self):
        with patch.object(providers,'ensure_ollama'),patch.object(providers,'request_json',return_value={'remote_host':'https://ollama.com'}):
            with self.assertRaises(ValueError):providers.complete('secret',{'provider':'ollama','model':'remote'},private=True)
    def test_mail_summary_always_local(self):
        services.save_settings({'provider':'openai','model':'gpt-test','cloud_private_data':True})
        with patch.object(services,'mail_read',return_value=[{'body':'confidentiel'}]),patch.object(services,'complete',return_value={'text':'résumé'}) as call,patch.object(services,'save_result',side_effect=lambda r,t:r):
            services.mail_summary({'account':'test'})
            self.assertEqual(call.call_args.args[1]['provider'],'ollama');self.assertFalse(call.call_args.kwargs['web'])
    def test_mail_in_briefing_always_local(self):
        services.save_settings({'provider':'openai','model':'gpt-test','include_mail_in_briefing':True,'cloud_private_data':True})
        with patch.object(services,'accounts',return_value={'test':{}}),patch.object(services,'mail_read',return_value=[{'body':'private'}]),patch.object(services,'complete',return_value={'text':'résultat'}) as call,patch.object(services,'save_result',side_effect=lambda r,t:r):
            services.briefing();self.assertEqual(call.call_args_list[1].args[1]['provider'],'ollama')
    def test_template_exports_replace_identity_and_keep_layout(self):
        for ident in templates.catalog()['items']:
            dest=self.home/(ident+'.docx');source=Document(templates.DIR/(ident+'.docx'))
            self.assertTrue(templates.export('Développement de test.\n\n1. Une question précise ?', 'Objet de test',dest,ident,{'recipient':'Madame la Ministre','salutation':'Madame la Ministre,'}))
            d=Document(dest);text='\n'.join(p.text for p in d.paragraphs)
            self.assertIn('Camille Exemple',text);self.assertIn('Madame la Ministre',text);self.assertNotIn('{{',text)
            self.assertEqual(d.sections[0].left_margin,source.sections[0].left_margin)
            self.assertEqual(len(d.inline_shapes),1)
            with ZipFile(dest) as z:
                self.assertEqual(len([n for n in z.namelist() if n.startswith('word/media/')]),1)
                xml=b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
                for word in [b'LauraHidalgo',b'GEOFFROY',b'Cauchie',b'Solarclick',b'Capitale europ']:
                    self.assertNotIn(word,xml)
    def test_index_and_query_remain_local(self):
        data={'name':'patrimoine.txt','content':base64.b64encode('Rénovation du patrimoine bruxellois'.encode()).decode()}
        services.index_documents({'files':[data]});matches=services.search_documents({'query':'patrimoine'})['matches']
        self.assertEqual(matches[0]['name'],'patrimoine.txt')
    def test_report_cannot_escape_reports_folder(self):
        for name in ['../settings.json','/etc/passwd','a.sh']:
            with self.assertRaises(ValueError):services.report_file(name)
    def test_eml_exact_bytes_and_one_use_confirmation(self):
        account={'type':'imap','address':'from@example.org'}
        raw=b'From: from@example.org\r\nTo: to@example.org\r\nSubject: Test\r\n\r\nContenu\r\n'
        data={'account':'test','file':{'name':'message.eml','content':base64.b64encode(raw).decode()}}
        api=desktop.Api();api._window=Mock();api._window.create_confirmation_dialog.return_value=True
        with patch.object(services,'accounts',return_value={'test':account}),patch.object(desktop,'smtp_send',return_value={'text':'accepted'}) as send:
            result=api.call('mail_preview',data)['data'];nonce=result['nonce']
            self.assertTrue(api.call('mail_send',{'nonce':nonce})['ok']);self.assertEqual(send.call_args.args[0][2],raw)
            self.assertFalse(api.call('mail_send',{'nonce':nonce})['ok']);self.assertEqual(send.call_count,1)
    def test_eml_sender_and_bcc_rejected(self):
        for raw in [b'From: someone@example.org\r\nTo: to@example.org\r\n\r\nTest',b'From: from@example.org\r\nTo: to@example.org\r\nBcc: hidden@example.org\r\n\r\nTest']:
            with patch.object(services,'accounts',return_value={'test':{'type':'imap','address':'from@example.org'}}):
                with self.assertRaises(ValueError):desktop.mail_preview({'account':'test','file':{'content':base64.b64encode(raw).decode()}})
    def test_cancel_send_does_not_send(self):
        api=desktop.Api();api._window=Mock();api._window.create_confirmation_dialog.return_value=False
        api._pending['nonce']=({},['to@example.org'],b'bytes',{'subject':'test'},'test')
        with patch.object(desktop,'smtp_send') as send:
            self.assertTrue(api.call('mail_send',{'nonce':'nonce'})['ok']);send.assert_not_called()
    def test_schedule_failure_restores_config_and_plist(self):
        path=self.home/'Library/LaunchAgents/net.compose.jarvis-parlement.briefing.plist';path.parent.mkdir(parents=True)
        original=plistlib.dumps({'Label':'old','StartCalendarInterval':{'Hour':10,'Minute':30}});path.write_bytes(original)
        common.atomic_json(self.cfg/'settings.json',{**common.settings(),'schedule_enabled':True})
        def launch(args,**kw):
            if kw.get('check'):raise RuntimeError('launchctl refused')
            return Mock(returncode=0)
        with patch.object(services.subprocess,'run',side_effect=launch):
            with self.assertRaises(ValueError):services.save_settings({'briefing_time':'11:00'})
        self.assertEqual(common.settings()['briefing_time'],'10:30');self.assertEqual(path.read_bytes(),original)
    def test_oauth_rejects_other_issuer(self):
        for url in ['http://auth.openai.com/jwks','https://auth.openai.com.evil.test/jwks','https://a@auth.openai.com/jwks']:
            with self.assertRaises(ValueError):chatgpt_auth._trusted_endpoint(url)
    def test_only_http_links_open(self):
        for url in ['file:///etc/passwd','javascript:alert(1)','https://user:pass@example.org']:
            with self.assertRaises(ValueError):desktop.safe_url(url)

if __name__=='__main__':unittest.main()
